from io import BytesIO
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.lib import colors
from xml.sax.saxutils import escape
import re

def markdown_to_lines(text):
    lines=[]
    for line in text.splitlines():
        s=line.strip()
        if s.startswith("### "): lines.append(("h2",s[4:]))
        elif s.startswith("## "): lines.append(("h1",s[3:]))
        elif s.startswith("# "): lines.append(("title",s[2:]))
        elif s.startswith("- ") or s.startswith("* "): lines.append(("bullet",s[2:]))
        elif re.match(r"^\d+\. ",s): lines.append(("bullet",s))
        elif s: lines.append(("p",s))
        else: lines.append(("space",""))
    return lines

def create_docx(title, content):
    doc=Document(); doc.add_heading(title,0)
    for kind,text in markdown_to_lines(content):
        if kind=="title": doc.add_heading(text,0)
        elif kind=="h1": doc.add_heading(text,1)
        elif kind=="h2": doc.add_heading(text,2)
        elif kind=="bullet": doc.add_paragraph(text,style="List Bullet")
        elif kind=="p": doc.add_paragraph(text)
    out=BytesIO(); doc.save(out); out.seek(0); return out

def create_pdf(title, content):
    out=BytesIO(); doc=SimpleDocTemplate(out,pagesize=A4,rightMargin=48,leftMargin=48,topMargin=52,bottomMargin=48)
    styles=getSampleStyleSheet(); styles.add(ParagraphStyle(name="SmallBody",parent=styles["BodyText"],fontSize=9.5,leading=14,spaceAfter=7,textColor=colors.HexColor("#273449")))
    story=[Paragraph(escape(title),styles["Title"]),Spacer(1,16)]
    for kind,text in markdown_to_lines(content):
        if kind=="space": story.append(Spacer(1,6)); continue
        safe=escape(text).replace("  "," &nbsp;")
        style=styles["Heading1"] if kind=="h1" else styles["Heading2"] if kind=="h2" else styles["Title"] if kind=="title" else styles["SmallBody"]
        if kind=="bullet": safe="• " + safe
        story.append(Paragraph(safe,style))
    doc.build(story); out.seek(0); return out
