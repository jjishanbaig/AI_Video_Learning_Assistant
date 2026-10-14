from pathlib import Path
import re
import fitz
from docx import Document
from pptx import Presentation
from youtube_transcript_api import YouTubeTranscriptApi

ALLOWED = {".pdf", ".docx", ".pptx", ".txt", ".md", ".mp4", ".mov", ".webm", ".m4v", ".mpeg", ".mpg"}
VIDEO_EXTS = {".mp4", ".mov", ".webm", ".m4v", ".mpeg", ".mpg"}
MIME = {".mp4": "video/mp4", ".mov": "video/quicktime", ".webm": "video/webm", ".m4v": "video/x-m4v", ".mpeg": "video/mpeg", ".mpg": "video/mpeg"}

def youtube_id(url: str):
    m = re.search(r"(?:youtu\.be/|youtube\.com/(?:watch\?v=|embed/|shorts/|live/))([\w-]{11})", url)
    if not m:
        raise ValueError("Please provide a valid YouTube video URL.")
    return m.group(1)

def fetch_youtube_transcript(url: str):
    vid = youtube_id(url)
    api = YouTubeTranscriptApi()
    try:
        transcript = api.fetch(vid)
    except Exception as e:
        try:
            transcripts = api.list(vid)
            transcript = transcripts.find_transcript(["en", "en-US", "en-GB"]).fetch()
        except Exception:
            raise ValueError("Could not retrieve a transcript for this video. Captions may be disabled or unavailable. Upload a permitted video file or paste the transcript instead.") from e
    lines = []
    for item in transcript:
        text = getattr(item, "text", None) or (item.get("text") if isinstance(item, dict) else "")
        start = getattr(item, "start", None) if not isinstance(item, dict) else item.get("start")
        if start is not None:
            s = int(float(start))
            stamp = f"[{s//3600:02d}:{(s%3600)//60:02d}:{s%60:02d}] "
        else:
            stamp = ""
        text = re.sub(r"\s+", " ", str(text)).strip()
        if text:
            lines.append(stamp + text)
    if not lines:
        raise ValueError("Transcript was empty. Try another video or upload/paste a transcript.")
    return "\n".join(lines), vid

def extract_document(path: Path):
    ext = path.suffix.lower()
    if ext == ".pdf":
        with fitz.open(path) as doc:
            return "\n\n".join(f"[Page {i+1}]\n{page.get_text()}" for i, page in enumerate(doc)).strip()
    if ext == ".docx":
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            parts.extend(" | ".join(c.text.strip() for c in row.cells) for row in table.rows)
        return "\n".join(parts).strip()
    if ext == ".pptx":
        prs = Presentation(path)
        parts = []
        for i, slide in enumerate(prs.slides, 1):
            texts = [sh.text for sh in slide.shapes if hasattr(sh, "text") and sh.text.strip()]
            if texts:
                parts.append(f"[Slide {i}]\n" + "\n".join(texts))
        return "\n\n".join(parts).strip()
    if ext in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="replace").strip()
    raise ValueError(f"Unsupported document type: {ext}")
